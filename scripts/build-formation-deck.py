"""Generates the training-progress PowerPoint deck for the Angular 21 training.

Run with: python scripts/build-formation-deck.py
Regenerate freely — this script is the source of truth, the .pptx is the output.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --- Palette (matches the tasks.store app) --------------------------------
INK = RGBColor(0x14, 0x17, 0x1C)
INK_SOFT = RGBColor(0x4B, 0x52, 0x59)
INK_FAINT = RGBColor(0x7C, 0x84, 0x8C)
PAPER = RGBColor(0xED, 0xF0, 0xF2)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xDC, 0xE1)
ACCENT = RGBColor(0xB8, 0x86, 0x2B)

SANS = "Segoe UI"
SANS_SEMIBOLD = "Segoe UI Semibold"
MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.75)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def add_rect(slide, x, y, w, h, color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box


def set_run(run, text, font=SANS, size=14, color=INK, bold=False, italic=False):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_header(slide, kicker, point_label=None):
    """Wordmark top-left ("tasks" in accent + ".store" in ink), progress label top-right."""
    brand = add_textbox(slide, MARGIN, Inches(0.45), Inches(4), Inches(0.4))
    p = brand.text_frame.paragraphs[0]
    set_run(p.add_run(), "tasks", font=MONO, size=15, color=ACCENT, bold=True)
    set_run(p.add_run(), ".store", font=MONO, size=15, color=INK, bold=True)

    if point_label:
        label = add_textbox(slide, SLIDE_W - MARGIN - Inches(2.5), Inches(0.45), Inches(2.5), Inches(0.4))
        label.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        set_run(label.text_frame.paragraphs[0].add_run(), point_label, font=SANS, size=12, color=INK_FAINT)

    kicker_box = add_textbox(slide, MARGIN, Inches(1.05), Inches(10), Inches(0.35))
    set_run(kicker_box.text_frame.paragraphs[0].add_run(), kicker, font=SANS, size=13, color=INK_SOFT)

    add_rect(slide, MARGIN, Inches(0.95), Inches(0.5), Pt(2.5), ACCENT)


def add_title(slide, title, y=Inches(1.4), size=32):
    box = add_textbox(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, Inches(1.1))
    set_run(box.text_frame.paragraphs[0].add_run(), title, font=SANS_SEMIBOLD, size=size, color=INK, bold=True)


def add_checklist(slide, items, x=MARGIN, y=Inches(2.6), w=None, line_gap=0.62):
    """Each item is a string; segments wrapped in backticks render in monospace."""
    if w is None:
        w = SLIDE_W - 2 * MARGIN
    box = add_textbox(slide, x, y, w, SLIDE_H - y - Inches(0.9))
    tf = box.text_frame
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(line_gap * 72 - 16)
        set_run(p.add_run(), "☐  ", font=SANS, size=16, color=ACCENT)
        parts = item.split("`")
        for i, part in enumerate(parts):
            if part == "":
                continue
            if i % 2 == 1:
                set_run(p.add_run(), part, font=MONO, size=14.5, color=INK)
            else:
                set_run(p.add_run(), part, font=SANS, size=16, color=INK)
    return box


def add_footer(slide, text):
    add_rect(slide, MARGIN, SLIDE_H - Inches(0.65), SLIDE_W - 2 * MARGIN, Pt(1), LINE)
    box = add_textbox(slide, MARGIN, SLIDE_H - Inches(0.55), SLIDE_W - 2 * MARGIN, Inches(0.4))
    p = box.text_frame.paragraphs[0]
    parts = text.split("`")
    for i, part in enumerate(parts):
        if part == "":
            continue
        if i % 2 == 1:
            set_run(p.add_run(), part, font=MONO, size=11.5, color=INK_SOFT)
        else:
            set_run(p.add_run(), part, font=SANS, size=11.5, color=INK_FAINT)


# --- Slide 1 — Title -------------------------------------------------------
slide = add_slide()
add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, ACCENT)
brand = add_textbox(slide, Inches(1.1), Inches(2.7), Inches(9), Inches(0.9))
p = brand.text_frame.paragraphs[0]
set_run(p.add_run(), "tasks", font=MONO, size=30, color=ACCENT, bold=True)
set_run(p.add_run(), ".store", font=MONO, size=30, color=INK, bold=True)

title = add_textbox(slide, Inches(1.1), Inches(3.5), Inches(10.5), Inches(1.3))
set_run(
    title.text_frame.paragraphs[0].add_run(),
    "Signals, détection, DI moderne & smart/dumb",
    font=SANS_SEMIBOLD,
    size=34,
    color=INK,
    bold=True,
)

sub = add_textbox(slide, Inches(1.1), Inches(4.45), Inches(10.5), Inches(0.6))
set_run(
    sub.text_frame.paragraphs[0].add_run(),
    "Formation Angular 21 — support de suivi",
    font=SANS,
    size=17,
    color=INK_SOFT,
)

# --- Slide 2 — Agenda -------------------------------------------------------
slide = add_slide()
add_header(slide, "Sommaire")
add_title(slide, "Ce que couvre cette formation")
add_checklist(
    slide,
    [
        "Mise en place du projet — Angular 21, standalone, zoneless",
        "Point 1 — Signals et opérateurs",
        "Point 2 — Stratégie de détection des composants",
        "Point 3 — Injection de dépendances moderne",
        "Point 4 — Pattern smart / dumb",
        "Suivi de progression",
    ],
)
add_footer(slide, "Application support : `tasks.store` — gestion de tâches")

# --- Slide 3 — Setup ---------------------------------------------------------
slide = add_slide()
add_header(slide, "Mise en place", "Setup")
add_title(slide, "Projet initial")
add_checklist(
    slide,
    [
        "`ng new` — Angular 21, standalone, zoneless",
        "Control flow natif partout — `@if` / `@for` / `@switch`, jamais `*ngIf` / `*ngFor`",
        "Structure par feature — `core/`, `shared/ui/`, `shared/data-access/`, `features/`",
        "Application support : gestion de tâches (liste, filtres, création, édition)",
        "Pas de backend réel — `TaskApiService` simule un appel avec `delay(300)`",
    ],
)
add_footer(slide, "`angular.json`, `tsconfig.json` (alias `@core`, `@shared`, `@features`)")

# --- Slide 4 — Point 1 -------------------------------------------------------
slide = add_slide()
add_header(slide, "Point 1 sur 4", "Signals")
add_title(slide, "Signals et opérateurs")
add_checklist(
    slide,
    [
        "`signal()` — état source : la liste des tâches",
        "`computed()` — état dérivé pur : compteur, liste filtrée",
        "`equal` personnalisé — comparaison de contenu sur un tableau/objet",
        "`linkedSignal()` — se réinitialise avec le contexte, reste modifiable ensuite",
        "`effect()` — effet de bord isolé, pas un `computed()` : persistance",
        "`effect()` + `untracked()` — lire sans dépendre, piège de boucle infinie évité",
        "`toSignal()` — pont vers un flux RxJS externe (lib tierce simulée)",
    ],
)
add_footer(slide, "`features/tasks/data-access/task-store.service.ts`")

# --- Slide 5 — Point 2 -------------------------------------------------------
slide = add_slide()
add_header(slide, "Point 2 sur 4", "Détection")
add_title(slide, "Stratégie de détection des composants")
add_checklist(
    slide,
    [
        "`ChangeDetectionStrategy.OnPush` partout — aucun `markForCheck()` par défaut",
        "Démo live : bug de mutation directe d'un tableau derrière un signal",
        "Correctif : `.update()` au lieu de `.push()` — nouvelle référence, vue notifiée",
        "Cas réel de `markForCheck()` — widget tiers hors radar Angular (`setInterval`)",
    ],
)
add_footer(slide, "`_exemples-a-ne-pas-suivre/buggy-task-list.example.ts`, `shared/ui/legacy-tick-counter/`")

# --- Slide 6 — Point 3 -------------------------------------------------------
slide = add_slide()
add_header(slide, "Point 3 sur 4", "Injection")
add_title(slide, "Injection de dépendances moderne")
add_checklist(
    slide,
    [
        "`inject()` en fonction partout — aucune injection par constructeur",
        "Service global `providedIn: 'root'` — `UserPreferencesService`",
        "Service scopé à la route — `TaskStore` fourni dans `tasks.routes.ts`",
        "L'état ne doit pas survivre à la navigation hors de la feature",
        "Guard fonctionnel `CanActivateFn` — pas de constructeur à injecter",
    ],
)
add_footer(slide, "`core/services/user-preferences.service.ts`, `core/guards/tasks-access.guard.ts`")

# --- Slide 7 — Point 4 -------------------------------------------------------
slide = add_slide()
add_header(slide, "Point 4 sur 4", "Smart / dumb")
add_title(slide, "Pattern smart / dumb")
add_checklist(
    slide,
    [
        "`TaskBoard` — composant smart : orchestre, ne présente pas",
        "Composants dumb — `input()` / `output()` uniquement : `TaskFilterBar`, `TaskCreateForm`",
        "`model()` — binding bidirectionnel simple sur la case « terminé »",
        "État UI local dans un dumb component — `isEditing` dans `TaskItem`",
        "Un dumb component peut porter de l'état tant que ce n'est pas de la donnée métier",
    ],
)
add_footer(slide, "`features/tasks/containers/task-board/`, `shared/ui/task-item/task-item.ts`")

# --- Slide 8 — Progress tracker ---------------------------------------------
slide = add_slide()
add_header(slide, "Suivi", "Récap")
add_title(slide, "Suivi de progression")

rows = [
    ("Setup", "Projet Angular 21 standalone / zoneless", ""),
    ("Point 1", "Signals et opérateurs", ""),
    ("Point 2", "Stratégie de détection des composants", ""),
    ("Point 3", "Injection de dépendances moderne", ""),
    ("Point 4", "Pattern smart / dumb", ""),
]
table_top = Inches(2.6)
table_w = SLIDE_W - 2 * MARGIN
table_h = Inches(3.6)
gfx = slide.shapes.add_table(len(rows) + 1, 3, MARGIN, table_top, table_w, table_h)
table = gfx.table
table.columns[0].width = Inches(1.6)
table.columns[1].width = table_w - Inches(1.6) - Inches(1.6)
table.columns[2].width = Inches(1.6)

headers = ["Étape", "Sujet", "Fait"]
for c, text in enumerate(headers):
    cell = table.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = INK
    cell.margin_left = Inches(0.15)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    set_run(tf.paragraphs[0].add_run(), text, font=SANS_SEMIBOLD, size=14, color=PAPER, bold=True)

for r, (step, topic, done) in enumerate(rows, start=1):
    for c, text in enumerate([step, topic, "☐"]):
        cell = table.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURFACE if r % 2 else PAPER
        cell.margin_left = Inches(0.15)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        font = MONO if c == 0 else SANS
        color = ACCENT if c == 2 else INK
        size = 20 if c == 2 else 14
        run = tf.paragraphs[0].add_run()
        if c == 2:
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        set_run(run, text, font=font, size=size, color=color)

add_footer(slide, "À cocher en direct pendant la session — voir aussi `README.md` du projet")

# --- Slide 9 — Closing --------------------------------------------------------
slide = add_slide()
add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, ACCENT)
box = add_textbox(slide, Inches(1.1), Inches(3.0), Inches(10.5), Inches(1.0))
set_run(box.text_frame.paragraphs[0].add_run(), "Merci — questions ?", font=SANS_SEMIBOLD, size=32, color=INK, bold=True)

sub = add_textbox(slide, Inches(1.1), Inches(3.85), Inches(10.5), Inches(1.1))
tf = sub.text_frame
p1 = tf.paragraphs[0]
set_run(p1.add_run(), "yarn start", font=MONO, size=15, color=ACCENT)
set_run(p1.add_run(), "  pour relancer la démo", font=SANS, size=15, color=INK_SOFT)
p2 = tf.add_paragraph()
set_run(p2.add_run(), "Détail fichier par fichier dans ", font=SANS, size=15, color=INK_SOFT)
set_run(p2.add_run(), "README.md", font=MONO, size=15, color=ACCENT)

prs.save("formation-angular21-signals.pptx")
print("Deck written to formation-angular21-signals.pptx")
